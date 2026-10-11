"""Tests for the CV module — adapter, validator, docx renderer."""

from __future__ import annotations

from html import escape
from pathlib import Path
from types import SimpleNamespace

import pytest

from smartapply.cv import (
    CvAdapter,
    CvDocxRenderer,
    CvValidator,
    HtmlApplicationRenderer,
)
from smartapply.cv.assembly import profile_experiences
from smartapply.cv.education import english_institution_name
from smartapply.cv.motivation_validator import (
    MotivationLetterValidator,
)
from smartapply.llm import (
    AdaptedBullet,
    AdaptedCV,
    ApplicationTexts,
    CVSelection,
    JobAnalysis,
    MockLLMProvider,
    MotivationLetter,
    MotivationLetterRepair,
)
from smartapply.llm.prompts import cv_adaptation
from smartapply.pipeline.apply.cv_writer import CvWriterMixin, _document_company_label
from smartapply.pipeline.reports import ApplyReport
from smartapply.profile import Degree, get_profile
from smartapply.utils.skill_terms import skill_term_matches

pytestmark = pytest.mark.usefixtures("isolated_db")


def _sample_analysis() -> JobAnalysis:
    return JobAnalysis(
        role_type="Data Scientist NLP",
        seniority="mid",
        domain="HealthTech",
        main_tasks=[
            "Build RAG pipelines",
            "Fine-tune transformer models",
            "Productionize ML services",
        ],
        required_skills=["PyTorch", "RAG", "Python"],
        nice_to_have=["Docker"],
        match_reasons=["Strong NLP background", "Multimodal AI experience"],
        risks=["No prior healthtech experience"],
        cv_keywords_to_include=["PyTorch", "RAG", "Whisper", "FAISS"],
    )




def _valid_adapted_cv() -> AdaptedCV:
    return AdaptedCV(
        cv_title="Data Scientist – NLP & Multimodal AI",
        professional_summary=(
            "Data Scientist with 2 years in NLP and machine learning. "
            "Built RAG and speech pipelines on synthetic datasets."
        ),
        selected_experiences=profile_experiences(get_profile()),
        selected_project_ids=["proj_evidence_rag", "proj_ner_camembert", "proj_smartapply"],
        skills_order=["ml_ai", "data_infra", "stats_signal"],
        warnings=[],
    )


def _empty_skills_cv() -> AdaptedCV:
    return AdaptedCV(
        cv_title="Adapted CV",
        professional_summary="Source-grounded summary.",
        selected_experiences=profile_experiences(get_profile()),
        selected_project_ids=[],
        skills_profile_id="mixed",
        selected_skills=[],
        skills_order=[],
        warnings=[],
    )


def _analysis_with_required_skills(required: list[str]) -> JobAnalysis:
    return JobAnalysis(
        role_type="Data Specialist",
        seniority="mid",
        domain="Data",
        main_tasks=[],
        required_skills=required,
        nice_to_have=[],
        match_reasons=[],
        risks=[],
        cv_keywords_to_include=[],
    )


def _selected_skill_map(cv: AdaptedCV) -> dict[str, list[str]]:
    return {block.category_id: list(block.skills) for block in cv.selected_skills}


# ---------------- Adapter ----------------


def test_offer_skill_matching_prefers_specific_nested_skills() -> None:
    profile = get_profile()
    adapter = CvAdapter(
        profile,
        llm=MockLLMProvider(),
        )

    cases = [
        ("PySpark", "data_infra", ["PySpark"]),
        ("experience with PySpark", "data_infra", ["PySpark"]),
        ("Spark", "data_infra", ["Spark"]),
        ("Mathematical Optimization", "stats_signal", ["Mathematical optimization"]),
        (
            "advanced mathematical optimization methods",
            "stats_signal",
            ["Mathematical optimization"],
        ),
        ("Optimization", "stats_signal", ["Optimization"]),
        ("Experience with PyTorch", "ml_ai", ["PyTorch"]),
    ]

    for term, category_id, expected_skills in cases:
        adapted = adapter._ensure_supported_offer_skills(
            _empty_skills_cv(),
            _analysis_with_required_skills([term]),
        )
        selected = _selected_skill_map(adapted)
        assert selected == {category_id: expected_skills}


@pytest.mark.parametrize(
    "term", ["théorie des graphes", "Connaissances en THEORIE DES GRAPHES", "Graph Theory"]
)
def test_graph_theory_translation_survives_cv_skill_filters(term: str) -> None:
    profile = get_profile().model_copy(deep=True)
    category = next(c for c in profile.skills.categories if c.id == "stats_signal")
    category.skills.append("Graph Theory")
    analysis = _analysis_with_required_skills([]).model_copy(
        update={"cv_keywords_to_include": [term], "role_type": "Optimization Engineer"}
    )
    adapter = CvAdapter(profile, llm=MockLLMProvider())

    adapted = adapter._ensure_supported_offer_skills(_empty_skills_cv(), analysis)
    assert _selected_skill_map(adapted) == {"stats_signal": ["Graph Theory"]}
    adapted = adapter._apply_role_family_contract(adapted, analysis, "Ingénieur Optimisation")
    adapted, removed = CvValidator(profile).auto_fix(adapted)
    assert "Graph Theory" in _selected_skill_map(adapted)["stats_signal"]
    assert not removed


def test_graph_theory_translation_does_not_infer_graph_tools_or_gnn_experience() -> None:
    assert not skill_term_matches("théorie des graphes", "Graph Neural Networks (GNNs)")
    for unsupported in ("NetworkX", "igraph", "graph-tool", "Graph Neural Networks (GNNs)"):
        assert not skill_term_matches(unsupported, "Graph Theory")
        assert (
            CvAdapter._matched_allowed_skill_keys(unsupported, {"graph theory": "Graph Theory"})
            == []
        )


def test_adapter_produces_letter_after_final_cv_and_review() -> None:
    from smartapply.llm import DocumentTextReview

    MockLLMProvider.clear()
    cv = _valid_adapted_cv()
    MockLLMProvider.register(
        "cv_adaptation",
        CVSelection(
            cv_title=cv.cv_title,
            selected_certificate_ids=[],
            selected_project_ids=cv.selected_project_ids,
            project_selection_reasons=[
                {
                    "project_id": project_id,
                    "reason": "Demonstrates documented NLP pipeline work.",
                    "offer_evidence": "Build RAG pipelines",
                    "complementary_value": "",
                }
                for project_id in cv.selected_project_ids
            ],
            skills_order=cv.skills_order,
            warnings=cv.warnings,
        ),
    )
    MockLLMProvider.register(
        "application_texts",
        ApplicationTexts(
            professional_summary=cv.professional_summary,
            motivation_letter_subject="Application: Data Scientist NLP",
            motivation_letter_body=_repair_valid_body(),
        ),
    )
    MockLLMProvider.register(
        "document_review",
        DocumentTextReview(
            issues=[],
        ),
    )
    adapter = CvAdapter(get_profile())
    adapted, letter = adapter.adapt_application(
        _sample_analysis(), job_title="Data Scientist NLP", job_company="Acme"
    )
    assert adapted.cv_title.startswith("Data Scientist")
    assert adapted.professional_summary == cv.professional_summary
    assert letter.subject.startswith("Application")
    assert adapted.selected_experiences
    assert adapted.selected_project_ids == cv.selected_project_ids


def test_adapter_keeps_three_final_projects_in_llm_order() -> None:
    profile = get_profile().model_copy(deep=True)
    if "Neo4j" not in profile.skills.allowed_skills:
        profile.skills.categories[0].skills.append("Neo4j")
    chosen = [project.id for project in reversed(profile.projects[:3])]
    personalization = CVSelection(
        cv_title="Data Scientist",
        selected_certificate_ids=[],
        selected_project_ids=chosen,
        project_selection_reasons=[
            {
                "project_id": project_id,
                "reason": "Demonstrates documented machine learning work.",
                "offer_evidence": "Productionize ML services",
                "complementary_value": "",
            }
            for project_id in chosen
        ],
        skills_order=[],
        warnings=[],
    )
    adapter = CvAdapter(
        profile,
        llm=MockLLMProvider({"cv_adaptation": personalization}),
        )

    adapted = adapter.adapt(
        _sample_analysis(), job_title="Data Scientist NLP", job_company="Example",
        job_description="Missions : construire un graphe de connaissances avec Neo4j.",
    )

    assert "professional_summary" not in adapted.model_dump()
    assert adapted.selected_project_ids == chosen
    assert "Neo4j" in {
        skill for block in adapted.selected_skills for skill in block.skills
    }


@pytest.mark.parametrize("available_count", [0, 1, 2, 5])
def test_html_keeps_empty_project_selection_regardless_of_catalogue(available_count: int) -> None:
    profile = get_profile().model_copy(
        update={"projects": get_profile().projects[:available_count]}
    )
    renderer = HtmlApplicationRenderer(profile)
    cv = _valid_adapted_cv().model_copy(update={"selected_project_ids": []})

    selected = renderer._select_projects(
        adapted=cv, project_by_id={project.id: project for project in profile.projects}
    )

    assert selected == []


# ---------------- Validator ----------------


def test_validator_accepts_clean_cv() -> None:
    v = CvValidator(get_profile())
    result = v.validate(_valid_adapted_cv())
    assert result.ok
    assert not result.errors


def test_validator_rejects_unknown_bullet_id() -> None:
    cv = _valid_adapted_cv()
    cv.selected_experiences[0].bullets[0] = AdaptedBullet(
        source_id="blt_does_not_exist",
        text="Did something cool.",
    )
    v = CvValidator(get_profile())
    result = v.validate(cv)
    assert not result.ok
    assert any("unknown_bullet_id" in e for e in result.errors)


def test_motivation_letter_validator_keeps_quality_checks_without_lexical_verdicts() -> None:
    analysis = _sample_analysis().model_copy(
        update={"required_skills": ["Kubernetes"], "cv_keywords_to_include": []}
    )
    letter = MotivationLetter(
        subject="Application: Data Scientist",
        body="My experience includes Kubernetes.",
    )
    result = MotivationLetterValidator(get_profile()).validate(
        letter,
        cv=_valid_adapted_cv(),
        analysis=analysis,
    )
    assert "letter_too_short:4" in result.warnings
    assert "letter_may_not_reference_selected_evidence" not in result.warnings
    assert not any(
        w.startswith(("unsupported_term_", "unsupported_tech_")) for w in result.warnings
    )


def test_cv_title_alignment_compares_role_families_not_offer_keywords() -> None:
    analysis = _sample_analysis().model_copy(
        update={
            "role_type": "Machine Learning Engineer",
            "main_tasks": ["Build RAG pipelines"],
            "required_skills": ["Python", "LangChain", "Kubernetes"],
            "cv_keywords_to_include": ["RAG", "vector search"],
        }
    )
    cv = _valid_adapted_cv().model_copy(
        update={
            "cv_title": "AI Engineer - Applied Intelligence",
            "professional_summary": "Machine learning profile for grounded retrieval workflows.",
        }
    )
    report = ApplyReport(job_id=42, application_id=None)

    CvWriterMixin()._validate_cv_offer_alignment(
        cv,
        analysis,
        "Machine Learning Engineer",
        report,
    )

    assert "cv_title_not_offer_anchored" not in report.validation_warnings


def test_cv_title_alignment_rejects_known_incompatible_family() -> None:
    analysis = _sample_analysis().model_copy(
        update={
            "role_type": "Data Engineer",
            "main_tasks": ["Build data pipelines"],
            "required_skills": ["Kubernetes"],
            "cv_keywords_to_include": [],
        }
    )
    cv = _valid_adapted_cv().model_copy(
        update={
            "cv_title": "Data Scientist - Kubernetes",
            "professional_summary": "Data engineering profile focused on reliable pipelines.",
        }
    )
    report = ApplyReport(job_id=42, application_id=None)

    CvWriterMixin()._validate_cv_offer_alignment(
        cv,
        analysis,
        "Data Engineer",
        report,
    )

    assert "cv_title_not_offer_anchored" in report.validation_warnings


def test_cv_title_alignment_ignores_ambiguous_title_family() -> None:
    analysis = _sample_analysis().model_copy(update={"role_type": "Data Engineer"})
    cv = _valid_adapted_cv().model_copy(update={"cv_title": "Applied Intelligence Specialist"})
    report = ApplyReport(job_id=42, application_id=None)

    CvWriterMixin()._validate_cv_offer_alignment(
        cv,
        analysis,
        "Data Engineer",
        report,
    )

    assert "cv_title_not_offer_anchored" not in report.validation_warnings


class _RepairRecordingProvider(MockLLMProvider):
    def __init__(self, registry=None):  # noqa: ANN001
        super().__init__(registry)
        self.calls: list[dict] = []

    def complete_json(self, **kwargs):  # noqa: ANN003, ANN201
        self.calls.append(kwargs)
        return super().complete_json(**kwargs)


class _RepairHarness(CvWriterMixin):
    def __init__(self, provider: MockLLMProvider):
        self.profile = get_profile()
        self.llm = provider
        self.letter_validator = MotivationLetterValidator(self.profile)


def _repair_valid_body() -> str:
    refrain = (
        "I value careful analysis, clear communication, reliable delivery, and "
        "thoughtful collaboration."
    )
    openings = (
        "The Evidence RAG Verifier project strengthened my approach to grounded "
        "retrieval, evidence review, and dependable evaluation.",
        "This selected work reflects a practical way of connecting technical "
        "decisions with clear product priorities and team needs.",
        "For this role, I would bring the same structured approach, attention "
        "to quality, and focus on useful outcomes.",
    )
    return "\n\n".join(f"{opening} {' '.join([refrain] * 4)}" for opening in openings)


def _repair_too_long_body() -> str:
    return f"{_repair_valid_body()} {' focused delivery' * 80}"


def _repair_once(
    harness: _RepairHarness,
    letter: MotivationLetter,
) -> MotivationLetter:
    return harness._repair_letter_once(
        letter,
        _valid_adapted_cv(),
        _sample_analysis(),
        job_title="Data Scientist",
        job_company="Acme",
        language="en",
        job_id=42,
    )


def test_single_letter_defect_is_repaired_once_with_smart_model() -> None:
    provider = _RepairRecordingProvider(
        {
            "motivation_letter_repair": MotivationLetterRepair(body=_repair_valid_body()),
        }
    )
    harness = _RepairHarness(provider)
    original = MotivationLetter(
        subject="Application - Data Scientist",
        body=_repair_too_long_body(),
    )

    repaired = _repair_once(harness, original)

    assert repaired.subject == original.subject
    assert repaired.body == _repair_valid_body()
    assert len(provider.calls) == 1
    call = provider.calls[0]
    assert call["purpose"] == "motivation_letter_repair"
    assert call["model"] == provider.smart_model
    assert call["temperature"] == 0.1
    assert call["schema"] is MotivationLetterRepair
    assert "letter_too_long" in call["user"]
    result = harness.letter_validator.validate(
        repaired,
        cv=_valid_adapted_cv(),
        analysis=_sample_analysis(),
    )
    assert result.errors == []
    assert result.warnings == []


def test_multiple_letter_defects_do_not_trigger_repair() -> None:
    provider = _RepairRecordingProvider()
    harness = _RepairHarness(provider)
    original = MotivationLetter(subject="Application", body="Brief note.")

    repaired = _repair_once(harness, original)

    assert repaired == original
    assert provider.calls == []


def test_unsuccessful_letter_repair_is_rejected_after_single_attempt() -> None:
    provider = _RepairRecordingProvider(
        {
            "motivation_letter_repair": MotivationLetterRepair(body="Still too brief."),
        }
    )
    harness = _RepairHarness(provider)
    original = MotivationLetter(subject="Application", body=_repair_too_long_body())

    repaired = _repair_once(harness, original)

    assert repaired == original
    assert len(provider.calls) == 1


def test_letter_repair_provider_failure_keeps_original() -> None:
    provider = _RepairRecordingProvider()
    harness = _RepairHarness(provider)
    original = MotivationLetter(subject="Application", body=_repair_too_long_body())

    repaired = _repair_once(harness, original)

    assert repaired == original
    assert len(provider.calls) == 1


# ---------------- DOCX renderer ----------------


def test_docx_renderer_produces_valid_file(tmp_path: Path) -> None:
    profile = get_profile()
    cv = _valid_adapted_cv()
    out = tmp_path / "cv.docx"
    renderer = CvDocxRenderer(profile)
    renderer.save(cv, out)
    assert out.exists()
    # File starts with the ZIP magic bytes (DOCX is a ZIP)
    assert out.read_bytes()[:2] == b"PK"
    # Re-open to confirm validity
    from docx import Document

    doc = Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Camille Martin" in full_text
    assert "Data Scientist" in full_text


def test_html_renderer_preserves_candidate_and_selected_project_content() -> None:
    profile = get_profile()
    cv = _valid_adapted_cv()
    html = HtmlApplicationRenderer(profile).render_cv_html(cv)

    for value in (profile.identity.full_name, cv.cv_title, cv.professional_summary):
        assert escape(value) in html
    for project in profile.projects:
        if project.id in cv.selected_project_ids:
            assert escape(project.name) in html
            if project.url:
                assert escape(str(project.url)) in html
        else:
            assert escape(project.name) not in html


def test_letter_renderer_uses_adapted_headline_and_fallback() -> None:
    profile = get_profile()
    renderer = HtmlApplicationRenderer(profile)
    letter = MotivationLetter(
        subject="Candidature - Data Scientist",
        body="Premier paragraphe.\n\nDeuxième paragraphe.\n\nTroisième paragraphe.",
    )

    html = renderer.render_letter_html(
        letter=letter,
        job_title="Data Scientist",
        job_company="Acme",
        letter_headline="Data Scientist - NLP & Multimodal AI",
    )

    assert '<p class="headline">Data Scientist - NLP &amp; Multimodal AI</p>' in html
    profile_headline = escape(profile.identity.title)
    assert f'<p class="headline">{profile_headline}</p>' not in html

    fallback_html = renderer.render_letter_html(
        letter=letter,
        job_title="Data Scientist",
        job_company="Acme",
    )

    assert f'<p class="headline">{profile_headline}</p>' in fallback_html


@pytest.mark.parametrize(
    ("job_company", "extracted_company_name", "language", "expected"),
    [
        ("Acme", "Acme", "fr", "Acme"),
        ("Acme", "", "fr", ""),
        ("Agence", "Entreprise non communiquée", "fr", ""),
        (
            "l'entreprise est inconnue ne pretends pas savoir le nom",
            "",
            "fr",
            "",
        ),
        ("Recruiter", "Unknown company", "en", ""),
    ],
)
def test_document_company_label_hides_anonymous_placeholders(
    job_company: str,
    extracted_company_name: str,
    language: str,
    expected: str,
) -> None:
    assert _document_company_label(job_company, extracted_company_name, language) == expected


@pytest.mark.parametrize("filename_candidate_name", [None, "Nour Lachtar"])
def test_application_renderer_passes_cv_title_to_letter_renderer(
    tmp_path: Path, filename_candidate_name: str | None
) -> None:
    from smartapply.pipeline.application_renderer import ApplicationDocumentRenderer
    from smartapply.pipeline.reports import ApplyReport

    profile = get_profile().model_copy(deep=True)
    profile.template_style.filename_candidate_name = filename_candidate_name
    original_full_name = profile.identity.full_name
    adapted = _valid_adapted_cv()
    original_cv_title = adapted.cv_title
    letter = MotivationLetter(
        subject="Candidature - Data Scientist",
        body="Premier paragraphe.\n\nDeuxième paragraphe.\n\nTroisième paragraphe.",
    )
    letter_calls: list[tuple[str, str | None]] = []
    cv_titles_seen: list[str] = []

    class FakeDocxRenderer:
        def save(self, adapted_cv, path):  # noqa: ANN001
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            Path(path).write_text("docx", encoding="utf-8")

    class FakeHtmlRenderer:
        def save_cv_html(self, adapted_cv, path):  # noqa: ANN001
            cv_titles_seen.append(adapted_cv.cv_title)
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            Path(path).write_text("cv html", encoding="utf-8")

        def save_cv_pdf(self, adapted_cv, path):  # noqa: ANN001
            cv_titles_seen.append(adapted_cv.cv_title)
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            Path(path).write_text("cv pdf", encoding="utf-8")

        def save_letter_html(self, **kwargs):  # noqa: ANN001
            letter_calls.append(("html", kwargs.get("letter_headline")))
            path = Path(kwargs["path"])
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("letter html", encoding="utf-8")

        def save_letter_pdf(self, **kwargs):  # noqa: ANN001
            letter_calls.append(("pdf", kwargs.get("letter_headline")))
            path = Path(kwargs["path"])
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("letter pdf", encoding="utf-8")

    renderer = ApplicationDocumentRenderer.__new__(ApplicationDocumentRenderer)
    renderer.profile = profile
    renderer.docx = FakeDocxRenderer()
    renderer.html = FakeHtmlRenderer()
    renderer.settings = SimpleNamespace(output_dir=tmp_path)
    report = ApplyReport(job_id=1, application_id=123)

    renderer.render_all(
        report=report,
        adapted=adapted,
        letter=letter,
        job_title="Data Scientist",
        job_company="Acme",
        language="fr",
    )

    assert letter_calls == [
        ("html", original_cv_title),
        ("pdf", original_cv_title),
    ]
    assert cv_titles_seen == [original_cv_title, original_cv_title]
    assert adapted.cv_title == original_cv_title
    assert profile.identity.full_name == original_full_name
    candidate_name = filename_candidate_name or original_full_name
    for field, prefix, extension in (
        ("docx_path", "CV", "docx"),
        ("cv_html_path", "CV", "html"),
        ("cv_pdf_path", "CV", "pdf"),
        ("letter_html_path", "Motivation", "html"),
        ("letter_pdf_path", "Motivation", "pdf"),
    ):
        path = Path(getattr(report, field))
        assert path.name == f"{prefix} - {candidate_name} - {original_cv_title}.{extension}"
        assert path.parent == tmp_path / "123"
        assert path.is_file()


def test_cv_renderers_translate_french_university_label(tmp_path: Path) -> None:
    assert english_institution_name("Sorbonne Université") == "Sorbonne University"

    profile = get_profile().model_copy(
        update={
            "education": [
                Degree(
                    id="edu_paris_cite",
                    title="Master's Degree",
                    field="Data Science",
                    institution="Université Paris Cité",
                    start_date="09/2022",
                    end_date="09/2024",
                    start_year=2022,
                    end_year=2024,
                )
            ]
        }
    )

    html = HtmlApplicationRenderer(profile).render_cv_html(_valid_adapted_cv())
    assert "Paris Cité University" in html
    assert "Université" not in html
    assert "09/2022 &ndash; 09/2024" in html

    out = tmp_path / "cv.docx"
    CvDocxRenderer(profile).save(_valid_adapted_cv(), out)

    from docx import Document

    doc = Document(str(out))
    full_text = "\n".join(
        p.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        for p in cell.paragraphs
    )
    assert "Paris Cité University" in full_text
    assert "Université" not in full_text
    assert "09/2022 – 09/2024" in full_text


def test_cv_prompt_separates_matching_keywords_from_display_skills() -> None:
    profile = get_profile()
    analysis = _sample_analysis().model_copy(
        update={
            "required_skills": [
                "Python",
                "Machine learning",
                "Computer vision",
                "MLOps",
                "Kubernetes",
                "Experience with PyTorch",
            ],
            "cv_keywords_to_include": ["PyTorch", "GCP"],
            "company_context": (
                "Acme Health builds patient-facing clinical AI products for care teams."
            ),
            "offer_interest_points": [
                "Work on RAG pipelines for medical knowledge access",
                "Collaborate with product and clinical teams",
            ],
        }
    )
    cv_prompt = cv_adaptation.build_user_prompt(
        profile=profile,
        analysis=analysis,
        job_title="Clinical AI Engineer",
        job_company="Acme Health",
        selected_experiences=profile.experiences[:1],
        selected_projects=profile.projects[:2],
    )
    assert "allowed_skills_by_category" in cv_prompt
    assert "core_skills_by_category" in cv_prompt
    assert "matching_keywords_for_profile_selection_not_display" in cv_prompt
    assert "project_selection_reasons" in cv_prompt
    profile_context, offer_context = cv_prompt.split("=== JOB ===", 1)
    for category in profile.skills.categories:
        for skill in category.skills:
            assert skill in profile_context
    for term in analysis.required_skills + analysis.cv_keywords_to_include:
        assert term in offer_context
