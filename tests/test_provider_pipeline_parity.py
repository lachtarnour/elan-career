"""Run real provider builders/readers through synthetic application workflows.

Only the final SDK request methods are replaced. No network or runtime data is
used; prompts, output schemas, cache, validation and publication stay production.
"""

from __future__ import annotations

import copy
import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from pydantic import BaseModel
from sqlalchemy import select

from smartapply.config import get_settings
from smartapply.cv.assembly import assemble_cv
from smartapply.database import session_scope
from smartapply.database.models import Application, Job, LLMCache, LLMUsage
from smartapply.database.session import init_db, reset_engine_cache
from smartapply.llm import (
    AdaptedCV,
    JobFitAssessment,
    MockLLMProvider,
    MotivationLetterRepair,
    OfferAnalysis,
    OfferExtractionOutput,
    get_llm_provider,
)
from smartapply.llm.anthropic_provider import AnthropicProvider
from smartapply.llm.openai_provider import OpenAIProvider
from smartapply.llm.prompts import motivation_letter_repair
from smartapply.llm.schemas import JobScopeScreening
from smartapply.offers import RawJob
from smartapply.pipeline import Pipeline
from smartapply.pipeline.ingestor import Ingestor
from smartapply.pipeline.process.analysis import _assemble_analysis
from smartapply.profile import get_profile
from tests.llm_config_helpers import configure_calls
from tests.test_integration import SAMPLE_JOBS, _register_llm_fixtures

_PURPOSE = {
    "JobScopeScreening": "job_scope_screening",
    "OfferExtractionOutput": "job_offer_extraction",
    "JobFitDecisionEvaluation": "job_fit_assessment",
    "CVSelection": "cv_adaptation",
    "ApplicationTexts": "application_texts",
    "DocumentTextReview": "document_review",
    "MotivationLetterRepair": "motivation_letter_repair",
}


def _responses():
    previous = dict(MockLLMProvider._default_registry)
    try:
        _register_llm_fixtures()
        fixtures = dict(MockLLMProvider._default_registry)
    finally:
        MockLLMProvider._default_registry = previous
    fixtures["job_offer_extraction"] = OfferExtractionOutput.model_validate(
        {
            **fixtures["job_offer_extraction"].model_dump(),
        }
    )
    fixtures["job_scope_screening"] = JobScopeScreening(
        mission_screen={"mission_relation": "direct", "mission_evidence": "sufficient"}
    )
    return fixtures


def _leaf_response(body, provider_name, fixtures, calls):
    if provider_name == "openai":
        schema_name = body["response_format"]["json_schema"]["name"]
        system, user = [message["content"] for message in body["messages"]]
        output_limit = body["max_completion_tokens"]
    else:
        schema_name = body["output_config"]["format"]["schema"]["title"]
        system = body["system"][0]["text"]
        user = body["messages"][0]["content"]
        output_limit = body["max_tokens"]
    purpose = _PURPOSE[schema_name]
    calls.append(
        {
            "purpose": purpose,
            "schema": schema_name,
            "system": system,
            "user": user,
            "model": body["model"],
            "temperature": body.get("temperature"),
            "effort": body.get("reasoning_effort", body.get("output_config", {}).get("effort")),
            "limit": output_limit,
        }
    )
    value = fixtures[purpose]
    if callable(value):
        value = value(user=user)
    if isinstance(value, BaseModel):
        value = value.model_dump(mode="json")
    content = json.dumps(value, ensure_ascii=False)
    if provider_name == "openai":
        return SimpleNamespace(
            model=body["model"],
            choices=[
                SimpleNamespace(
                    finish_reason="stop",
                    message=SimpleNamespace(content=content, refusal=None),
                )
            ],
            usage=SimpleNamespace(
                prompt_tokens=1300,
                completion_tokens=300,
                prompt_tokens_details=SimpleNamespace(cached_tokens=100, cache_write_tokens=200),
            ),
        )
    return {
        "model": body["model"],
        "stop_reason": "end_turn",
        "content": [{"type": "text", "text": content}],
        "usage": {
            "input_tokens": 1000,
            "cache_read_input_tokens": 100,
            "cache_creation_input_tokens": 200,
            "output_tokens": 300,
        },
    }


def _mock_leaf(monkeypatch, fixtures):
    calls = {"openai": [], "anthropic": []}

    def openai_call(self, **body):
        return _leaf_response(body, "openai", fixtures, calls["openai"])

    def anthropic_call(self, **body):
        return _leaf_response(body, "anthropic", fixtures, calls["anthropic"])

    monkeypatch.setattr(OpenAIProvider, "_call_openai", openai_call)
    monkeypatch.setattr(AnthropicProvider, "_call_anthropic", anthropic_call)
    return calls


def _switch_test_runtime(monkeypatch, directory, provider_name):
    directory.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{directory / 'database.sqlite'}")
    monkeypatch.setenv("OUTPUT_DIR", str(directory / "documents"))
    cheap = "gpt-5.4-mini" if provider_name == "openai" else "claude-haiku-5-5"
    smart = "gpt-5.4" if provider_name == "openai" else "claude-sonnet-5-5"
    configure_calls(monkeypatch, directory, **{
        purpose: {
            "provider": provider_name,
            "model": cheap if purpose in {
                "job_scope_screening", "job_offer_extraction", "job_fit_assessment"
            } else smart,
            "reasoning": "medium" if purpose in {"cv_adaptation", "application_texts"} else "high",
            "max_output_tokens": {
                "job_scope_screening": 4096, "job_offer_extraction": 6000,
                "job_fit_assessment": 12000,
            }.get(purpose, 6000),
        }
        for purpose in _PURPOSE.values() if purpose != "motivation_letter_repair"
    })
    get_settings.cache_clear()
    reset_engine_cache()
    init_db(run_maintenance=False)


def _prompt_contract(call):
    # Providers have different native generation parameters, while business
    # prompts and schemas stay byte-identical throughout the real pipeline.
    return {
        key: value
        for key, value in call.items()
        if key not in {"model", "temperature", "effort", "limit"}
    }


def test_both_real_providers_publish_the_same_nlp_application(isolated_db, monkeypatch):
    fixtures = _responses()
    calls = _mock_leaf(monkeypatch, fixtures)
    snapshots = {}
    for provider_name in ("openai", "anthropic"):
        _switch_test_runtime(monkeypatch, isolated_db / provider_name, provider_name)
        pipeline = Pipeline()
        assert isinstance(
            pipeline.llm.for_purpose("job_offer_extraction"),
            OpenAIProvider if provider_name == "openai" else AnthropicProvider
        )
        ingest = Ingestor()._persist(
            "fixture_search",
            [
                RawJob(
                    external_id="provider-parity:nlp",
                    source="fixture_search",
                    **SAMPLE_JOBS[0],
                )
            ],
        )
        assert ingest.inserted == 1
        processed = pipeline.process_pending()
        assert processed.total == 1 and processed.analyzed == 1
        report = pipeline.apply_to(ingest.job_ids[0])
        assert report.application_id is not None and not report.validation_errors
        for output in (report.docx_path, report.cv_pdf_path, report.letter_pdf_path):
            assert output and Path(output).is_file()
        docx = Document(report.docx_path)
        text = "\n".join(paragraph.text for paragraph in docx.paragraphs)
        text += "\n".join(
            cell.text for table in docx.tables for row in table.rows for cell in row.cells
        )
        assert all(fact in text for fact in ("Camille Martin", "Aurore Labs", "PyTorch", "0.82"))
        with session_scope() as session:
            application = session.get(Application, report.application_id)
            job = session.get(Job, ingest.job_ids[0])
            assert application.status == "ready_for_form_submission"
            assert len(application.documents) >= 8
            assert application.cv_json["selected_project_ids"] == [
                "proj_evidence_rag",
                "proj_ner_camembert",
                "proj_gpt2",
            ]
            assert job.analysis.fit_score == 0.75
            assert (
                job.analysis.fit_diagnostics["context"].get("provider", "openai") == provider_name
            )
            usage = list(session.scalars(select(LLMUsage).order_by(LLMUsage.id)))
            assert [entry.purpose for entry in usage] == list(_PURPOSE.values())[:-1]
            assert all(entry.provider == provider_name and not entry.cached for entry in usage)
            assert all(
                (entry.prompt_tokens, entry.completion_tokens) == (1300, 300) for entry in usage
            )
            cache = list(session.scalars(select(LLMCache)))
            assert len(cache) == 5 and all(entry.provider == provider_name for entry in cache)
            letter = next(
                doc.content for doc in application.documents if doc.doc_type == "motivation_letter"
            )
            snapshots[provider_name] = {
                "cv": copy.deepcopy(application.cv_json),
                "letter": letter,
                "analysis": copy.deepcopy(job.analysis.raw_response),
                "review_issues": [
                    attempt["review"]["issues"] for attempt in report.document_review
                ],
            }
        # An identical extraction uses the actual provider cache without another
        # API call, including after the complete application has been published.
        extraction = calls[provider_name][1]
        before = len(calls[provider_name])
        cached = pipeline.llm.complete_json(
            system=extraction["system"],
            user=extraction["user"],
            schema=OfferExtractionOutput,
            model=pipeline.llm.cheap_model,
            purpose="job_offer_extraction",
            job_id=ingest.job_ids[0],
        )
        assert cached.main_tasks == fixtures["job_offer_extraction"].main_tasks
        assert len(calls[provider_name]) == before
        with session_scope() as session:
            entry = session.scalar(select(LLMUsage).order_by(LLMUsage.id.desc()))
            assert entry.provider == provider_name and entry.cached and entry.cost_usd == 0
        assert [call["purpose"] for call in calls[provider_name]] == list(_PURPOSE.values())[:-1]
        assert [call["model"] for call in calls[provider_name]] == [
            pipeline.llm.cheap_model,
            pipeline.llm.cheap_model,
            pipeline.llm.cheap_model,
            pipeline.llm.smart_model,
            pipeline.llm.smart_model,
            pipeline.llm.smart_model,
        ]
        if provider_name == "anthropic":
            assert [call["effort"] for call in calls[provider_name]] == [
                "high",
                "high",
                "high",
                "medium",
                "medium",
                "high",
            ]
            assert all(call["temperature"] is None for call in calls[provider_name])
            assert [call["limit"] for call in calls[provider_name]] == [
                4096, 6000, 12000, 6000, 6000, 6000,
            ]
        else:
            assert [call["effort"] for call in calls[provider_name]] == [
                "high",
                "high",
                "high",
                "medium",
                "medium",
                "high",
            ]
            assert all(call["temperature"] is None for call in calls[provider_name])
    assert snapshots["openai"] == snapshots["anthropic"]
    assert [_prompt_contract(call) for call in calls["openai"]] == [
        _prompt_contract(call) for call in calls["anthropic"]
    ]


def test_letter_repair_contract_keeps_the_same_prompts_and_outputs(isolated_db, monkeypatch):
    stage = "motivation_letter_repair"
    fixtures = _responses()
    profile = get_profile()
    analysis = _assemble_analysis(
        OfferAnalysis.model_validate(
            fixtures["job_offer_extraction"].model_dump(exclude={"mission_screen"})
        ),
        JobFitAssessment(fit_score=0.59, match_reasons=["Synthetic NLP match"], risks=[]),
    )
    cv = AdaptedCV.model_validate(
        {
            **assemble_cv(fixtures["cv_adaptation"], profile).model_dump(),
            "professional_summary": fixtures["application_texts"].professional_summary,
        }
    )
    letter = fixtures["application_texts"].to_motivation_letter()
    system = motivation_letter_repair.build_system_prompt()
    user = motivation_letter_repair.build_user_prompt(
        profile=profile,
        analysis=analysis,
        adapted_cv=cv,
        letter=letter,
        job_title=SAMPLE_JOBS[0]["title"],
        job_company=SAMPLE_JOBS[0]["company"],
        language="fr",
        defect="Synthetic punctuation defect",
    )
    schema, fixtures[stage] = MotivationLetterRepair, MotivationLetterRepair(body=letter.body)
    calls = _mock_leaf(monkeypatch, fixtures)
    outputs = {}
    for provider_name in ("openai", "anthropic"):
        _switch_test_runtime(monkeypatch, isolated_db / (provider_name + "-repair"), provider_name)
        model = "gpt-5.4" if provider_name == "openai" else "claude-sonnet-5-5"
        configure_calls(monkeypatch, isolated_db / (provider_name + "-repair"), **{
            stage: {"provider": provider_name, "model": model, "reasoning": "high",
                    "max_output_tokens": 6000}
        })
        provider = get_llm_provider()
        outputs[provider_name] = provider.complete_json(
            system=system,
            user=user,
            schema=schema,
            model=provider.smart_model,
            purpose=stage,
        ).model_dump()
        assert len(calls[provider_name]) == 1
        assert calls[provider_name][0]["model"] == provider.smart_model
        assert calls[provider_name][0]["effort"] == "high"
        assert calls[provider_name][0]["temperature"] is None
    assert outputs["openai"] == outputs["anthropic"]
    assert _prompt_contract(calls["openai"][0]) == _prompt_contract(calls["anthropic"][0])
